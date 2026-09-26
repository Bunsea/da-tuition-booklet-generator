import unittest
import os
from pptx import Presentation
from pptx.util import Mm
import build_feedback_pptx

class TestFeedbackPptx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Generate the PPTX if not present
        if not os.path.exists(build_feedback_pptx.PPTX_OUTPUT_PATH):
            build_feedback_pptx.generate_powerpoint()

    def test_pptx_file_exists_and_valid(self):
        self.assertTrue(os.path.exists(build_feedback_pptx.PPTX_OUTPUT_PATH))
        prs = Presentation(build_feedback_pptx.PPTX_OUTPUT_PATH)
        self.assertEqual(len(prs.slides), 2, "PowerPoint presentation must have exactly 2 slides")

    def test_a4_portrait_dimensions(self):
        prs = Presentation(build_feedback_pptx.PPTX_OUTPUT_PATH)
        # 210mm x 297mm
        self.assertEqual(prs.slide_width, Mm(210))
        self.assertEqual(prs.slide_height, Mm(297))

    def test_slide1_content(self):
        prs = Presentation(build_feedback_pptx.PPTX_OUTPUT_PATH)
        slide1 = prs.slides[0]

        slide1_text = " ".join([shape.text for shape in slide1.shapes if shape.has_text_frame])
        
        # Check branding and title
        self.assertIn("DA TUITION MATHEMATICS", slide1_text)
        self.assertIn("Teacher Lesson & Progress Log", slide1_text)
        self.assertIn("Student:", slide1_text)
        self.assertIn("Tutor:", slide1_text)

        # Check all 6 teacher sections
        self.assertIn("[01]", slide1_text)
        self.assertIn("Topics & Concepts Taught Today", slide1_text)
        self.assertIn("What We Covered Today", slide1_text)
        self.assertIn("[02]", slide1_text)
        self.assertIn("Student Observations", slide1_text)
        self.assertIn("Key Strengths", slide1_text)
        self.assertIn("Areas to Work On", slide1_text)
        self.assertIn("[03]", slide1_text)
        self.assertIn("Learning Habits, Recall & Homework Check", slide1_text)
        self.assertIn("Prior Recall", slide1_text)
        self.assertIn("Homework Done", slide1_text)
        self.assertIn("[04]", slide1_text)
        self.assertIn("Classroom Behaviour & Working Speed", slide1_text)
        self.assertIn("[05]", slide1_text)
        self.assertIn("Any Concerns & Areas Needing Help", slide1_text)
        self.assertIn("[06]", slide1_text)
        self.assertIn("IN-CLASS QUIZ / PRACTICE SCORE", slide1_text)
        self.assertIn("Next Lesson Plan", slide1_text)

        # Strictly assert no academic jargon
        self.assertNotIn("Pedagogical", slide1_text)
        self.assertNotIn("Qualitative Diagnostic", slide1_text)
        self.assertNotIn("Prescription & Targets", slide1_text)

    def test_slide2_content(self):
        prs = Presentation(build_feedback_pptx.PPTX_OUTPUT_PATH)
        slide2 = prs.slides[1]

        slide2_text = " ".join([shape.text for shape in slide2.shapes if shape.has_text_frame])

        # Check student voice branding
        self.assertIn("Student Voice & Parent Partnership", slide2_text)
        self.assertIn("Student Voice Section", slide2_text)

        # Check Student Voice modules
        self.assertIn("[01]", slide2_text)
        self.assertIn("How Was Today's Lesson?", slide2_text)
        self.assertIn("Lesson Pacing", slide2_text)
        self.assertIn("Tutor Explanations", slide2_text)
        self.assertIn("[02]", slide2_text)
        self.assertIn("How Did You Feel About Your Tutor Today?", slide2_text)
        self.assertIn("Overall Tutor Rating", slide2_text)
        self.assertIn("[03]", slide2_text)
        self.assertIn("Student Requests & Change Hotline", slide2_text)
        self.assertIn("Request Change of Teacher", slide2_text)
        self.assertIn("Confidence Meter", slide2_text)

        # Check Tear-off line
        self.assertIn("✂", slide2_text)
        self.assertIn("tear off", slide2_text.lower())

        # Check Parent Partnership modules
        self.assertIn("[04]", slide2_text)
        self.assertIn("Parent Section & Home Study Update", slide2_text)
        self.assertIn("Home Study Verification", slide2_text)
        self.assertIn("Tutor's Note to Parents", slide2_text)
        self.assertIn("Assessment Alert", slide2_text)
        self.assertIn("Parent Notes", slide2_text)
        self.assertIn("Parent Signature", slide2_text)

        # Strictly assert no academic jargon
        self.assertNotIn("Pedagogical", slide2_text)

if __name__ == "__main__":
    unittest.main()
